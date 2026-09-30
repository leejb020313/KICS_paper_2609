# -*- coding: utf-8 -*-
"""Build the 2-page KICS-format draft of the JEV paper (paper_jev/jev_kics_draft.docx): JEV -> LLM fused cascade.

Every number is read from results/jev_strict_full.json (scripts/jev_strict.py: same data, splits, 5 learning-set
subsamples and McNemar-per-run protocol as the main paper; call rates chosen on the learning set), data sizes from
results/final_results_full.json. Figure 1 is figures/jev_method_diagram.pdf (TikZ), Figure 2 figures/jev_budget_full.png
(scripts/make_jev_figure.py). Math is LaTeX rendered to PNG by equations/render_equations.py. Layout code is shared with
paper/build_docx.py (KICS Word template, house style of the peer KICS papers).

    uv run --isolated --no-project --with python-docx --with pymupdf python paper_jev/build_docx.py

Convert to PDF with Word, then check it:
    uv run --locked --with pymupdf python scripts/verify_jev_paper_numbers.py paper_jev/jev_kics_draft.pdf
    uv run --no-project --with pymupdf python paper_jev/references/verify_references.py paper_jev/jev_kics_draft.pdf --online
"""
import json
import math
import os
import re

import pymupdf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
R = json.load(open(os.path.join(ROOT, "results", "jev_strict_full.json"), encoding="utf-8"))
R_MAIN = json.load(open(os.path.join(ROOT, "results", "final_results_full.json"), encoding="utf-8"))  # data sizes
MAJ = 3  # "5회 중 3회 이상 유의" = significant on a majority of the runs


def test(ds, claim):
    return R[ds]["tests"][claim]


def nsig(ds, claim):
    return test(ds, claim)["n_sig"]


def sign(ds, claim):
    """+1 / -1 when every run's difference has that sign, else 0."""
    ds_ = test(ds, claim)["diff"]
    return 1 if all(d > 0 for d in ds_) else -1 if all(d < 0 for d in ds_) else 0


def rate(ds, v="jevfuse"):
    return R[ds]["rate"][v]


def pct(b):
    return f"{round(100 * b)}%"


def sig_range(ds):
    """Call rates (> 0) at which fusion beats replacement on a majority of runs, as "20~60%" or "30% 이상"."""
    bs = (.05, .1, .2, .3, .4, .5, .6, .7, .8, .9, 1.0)
    rates = [b for b in bs if nsig(ds, f"jevfuse_{b}_vs_jevrep_{b}") >= MAJ]
    lo, hi = rates[0], rates[-1]
    assert [b for b in bs if lo <= b <= hi] == rates, "significant call rates are not contiguous"
    return f"{pct(lo)} 이상" if hi == 1.0 else f"{pct(lo)[:-1]}~{pct(hi)}"

BODY_FONT = "HY신명조"  # the KICS template's 한양신명조 (installed as HY신명조, H2MJSM.TTF)
# layout of the official KICS Word template (paper/template/kics_sample_word.doc, conf.kics.or.kr/2026f):
# title block 1.5 cm margins; two-column body T/B/L 2 cm, R 1.5 cm, column gap 0.75 cm; body and references 9 pt


def acc(ds, name):
    return R[ds]["metrics"][name]["acc"][0]


def f3(x):
    return f"{x:.3f}"


def ro(num):
    """Korean instrumental particle after a number read digit-by-digit: 으로 after a final 0/3/6 (영/삼/육), else 로."""
    return f"{num}으로" if num[-1] in "036" else f"{num}로"


def eul(num):
    """Korean object particle after a number read digit-by-digit: 을 after a final 0/1/3/6/7/8 (batchim), else 를."""
    return f"{num}을" if num[-1] in "013678" else f"{num}를"


def pp(ds, a, b):
    """Accuracy difference a - b in percentage points, one decimal."""
    return f"{100 * (acc(ds, a) - acc(ds, b)):.1f}"


def cal(ds, name):
    return R[ds]["calib_curve"][name][0]


def pval(ds, key):
    return R[ds]["contrasts"][key]["mcnemar_p"]


def pfmt(p):
    """McNemar p as printed: p<0.001 below that, three decimals below 0.1, otherwise two."""
    return "p<0.001" if p < 0.001 else f"p={p:.3f}" if p < 0.1 else f"p={p:.2f}"


def best_full(ds):
    """The stronger of the two full-call settings on this dataset (raw 0.5 cut-off or calib-tuned threshold)."""
    return max(("sonnet_raw", "sonnet_thr"), key=lambda n: acc(ds, n))


LATIN_FONT = "Times New Roman"  # Latin letters and digits; HY신명조's hyphen is drawn as a long dash


def set_font(run, name=BODY_FONT, size=9.5, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = LATIN_FONT if name == BODY_FONT else name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def centered(doc, text, size=10, bold=False, font=BODY_FONT, after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(text), font, size, bold)
    return p


def heading(doc, text, size=9):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4.75)
    p.paragraph_format.keep_with_next = True  # never leave a heading alone at the bottom of a column
    set_font(p.add_run(text), BODY_FONT, size, bold=True)


def subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), BODY_FONT, 9, bold=True)


def figure(doc, png, width_cm, caption, detail=""):
    """Figure with a bold caption lead (number + the finding) and an optional plain-text detail."""
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(3)
    pic.paragraph_format.space_after = Pt(1)
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(png, width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if detail else WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(5)
    set_font(cap.add_run(caption), BODY_FONT, 8, bold=True)
    if detail:
        set_font(cap.add_run(" " + detail), BODY_FONT, 8)


BODY_LINE_SPACING = 1.1  # 1.0 read as cramped at 9 pt in two columns


def body(doc, text, size=9, indent=0.3175, after=3.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, lead=None):
    """Body paragraph; `lead` is a bold run-in sentence that states the paragraph's finding."""
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.line_spacing = BODY_LINE_SPACING
    if indent:
        pf.first_line_indent = Cm(indent)
    if lead:
        set_font(p.add_run(lead + " "), BODY_FONT, size, bold=True)
    set_font(p.add_run(text), BODY_FONT, size)
    return p


def add_math(run, name, scale=0.9):
    """Insert equations/<name>.png at the natural size of its LaTeX PDF (10pt math scaled to the 9pt body)."""
    rect = pymupdf.open(os.path.join(HERE, "equations", name + ".pdf"))[0].rect
    run.add_picture(os.path.join(HERE, "equations", name + ".png"), width=Pt(rect.width * scale))
    # w:position is in half-points: drop the image by the depth below the LaTeX baseline (~1/4 of its height)
    run._element.get_or_add_rPr().append(OxmlElement("w:position"))
    run._element.rPr[-1].set(qn("w:val"), str(-round(rect.height * scale * 0.25 * 2)))


def columns(doc, num):
    sectPr = doc.sections[-1]._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(0.75 * 567)))


doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(1.5)

# ---------------- title block ----------------
centered(doc, "제로샷 결정 모델과 대형 언어 모델의 선택적 결합을 통한", 13, True, after=0)
centered(doc, "팩트체크 필요성 탐지", 13, True, after=0)
# title block as in the peer KICS papers: authors with the corresponding author starred, affiliation, e-mails
AUTHORS_KO, AUTHORS_EN = "이정빈, 김은경*", "Lee Jeongbin, Kim Eunkyung*"  # English romanization of the advisor: to confirm
AFFIL_KO, AFFIL_EN = "국립한밭대학교", "Hanbat National Univ."
EMAILS = "leejb020313@gmail.com, *(교신저자 이메일)"  # corresponding author's e-mail: to fill in
centered(doc, AUTHORS_KO, 11, after=0).paragraph_format.space_before = Pt(7.6)
centered(doc, AFFIL_KO, 11, after=0)
centered(doc, EMAILS, 11, after=7.6)
centered(doc, "Check-Worthy Claim Detection by Selectively Fusing\na Zero-Shot Decision Model and a Large Language Model",
         13, True, after=0)
centered(doc, AUTHORS_EN, 11, after=0).paragraph_format.space_before = Pt(14.2)
centered(doc, AFFIL_EN, 11, after=14.2)

cb, cl = "cb", "clef"
BEST = {d: R[d]["best_full"] for d in (cl, cb)}
FSEL = {d: f"jevfuse_{rate(d)}" for d in (cl, cb)}          # the proposed row
SSEL = {d: f"fuse_{rate(d, 'fuse')}" for d in (cl, cb)}      # SVM -> LLM fused at its own learning-set rate
RSEL = {d: f"jevrep_{rate(d)}" for d in (cl, cb)}           # JEV -> LLM replacement at the proposed rate
gain = {d: 100 * (acc(d, "jev_thr") - acc(d, "svm")) for d in (cl, cb)}
# directions the text asserts
assert rate(cb) == 0 and rate(cl) > 0  # "ClaimBuster에서는 LLM 호출을 선택하지 않았다"
assert all(nsig(d, "jev_thr_vs_svm") == 5 and sign(d, "jev_thr_vs_svm") > 0 for d in (cl, cb))
assert sign(cl, "jev_thr_vs_best_full") < 0 and sign(cb, "jev_thr_vs_best_full") > 0 and nsig(cb, "jev_thr_vs_best_full") == 5
assert nsig(cl, "jevfuse_sel_vs_best_full") == 0 and nsig(cl, "table:fuse_0.5") == 0 and SSEL[cl] == "fuse_0.5"
assert nsig(cb, f"table:{SSEL[cb]}") == 0
assert all(nsig(d, "jevfuse_sel_vs_sonnet_nnppi") == 5 and sign(d, "jevfuse_sel_vs_sonnet_nnppi") > 0 for d in (cl, cb))
assert all(acc(d, f"jevfuse_{b}") >= acc(d, f"jevrep_{b}") for d in (cl, cb) for b in (.05, .1, .2, .3, .4, .5, .6, .7, .8, .9, 1.0))
CB_PEAK = max((.05, .1, .2, .3), key=lambda b: acc(cb, f"jevfuse_{b}"))  # the test-curve peak the learning set did not pick

centered(doc, "요 약", 11, False, after=3)
body(doc, (
    "팩트체크 필요성 탐지(check-worthiness detection)는 유입되는 문장 중 사실 확인이 필요한 문장을 선별하는 기술로, "
    "모든 문장에 대형 언어 모델(LLM)을 호출하면 비용과 지연이 크다. 본 논문에서는 텍스트 대신 라벨 확률을 반환하는 결정 전용 "
    "모델 JEV를 1단계로 두고, JEV가 불확실하게 판정한 문장에만 LLM을 호출하여 두 모델의 점수를 결합하는 캐스케이드를 제안한다. "
    f"CLEF 2024와 ClaimBuster에서 평가한 결과, JEV는 라벨로 학습한 임베딩 분류기보다 정확도가 {gain[cl]:.1f}%p, "
    f"{gain[cb]:.1f}%p 높았다. 학습 세트에서 정한 LLM 호출률은 CLEF {pct(rate(cl))}, ClaimBuster {pct(rate(cb))}였으며, "
    f"이때 정확도({f3(acc(cl, FSEL[cl]))}, {f3(acc(cb, FSEL[cb]))})는 CLEF에서 LLM 전량 호출과 유의한 차이가 없었고 "
    "ClaimBuster에서는 유의하게 높았다. 또한 두 점수를 결합하는 방식은 LLM을 호출하는 모든 호출률에서 LLM의 판정으로 "
    "교체하는 방식보다 평균 정확도가 높았다."
), size=9, indent=0.5, after=2.5)

doc.add_section(0)
_body = doc.sections[-1]
_body.top_margin = _body.bottom_margin = _body.left_margin = Cm(2.0)
_body.right_margin = Cm(1.5)
columns(doc, 2)

# ---------------- I. 서론 ----------------
heading(doc, "Ⅰ. 서 론")
body(doc, (
    "팩트체크 필요성 탐지는 유입되는 모든 문장을 대상으로 사실 확인이 필요한지를 판정하므로 처리할 문장의 수가 많다. "
    "실제 팩트체크 서비스 운영사는 모든 문장에 LLM을 적용하는 것이 비현실적이라고 보고하였으며[1], 이에 저비용 모델이 먼저 "
    "판정하고 불확실한 문장만 LLM에 넘기는 캐스케이드가 대안이 된다. 최근 공개된 JEV는 텍스트를 생성하지 않고 라벨마다 확률을 "
    "반환하는 결정 전용 모델로, 생성형 LLM보다 훨씬 저렴하다(1,000건당 0.044달러)[2]."
))
body(doc, (
    "JEV를 1단계로 한 캐스케이드는 평가(judging) 과제에서 연구되었으나, JEV가 확신할 때는 그 판정을 받아들이고 나머지는 "
    "LLM의 판정으로 교체하는 방식이었다[2]. 또한 JEV와 LLM은 같은 곳에서 틀리는 경향이 있어, 이러한 캐스케이드는 비용만 "
    "줄일 뿐 정확도는 거의 높이지 못한다고 보고되었다[3]. 두 모델의 출력을 결합하는 방식은 향후 과제로 남아 있다[2]. 본 "
    "논문에서는 JEV를 팩트체크 필요성 탐지의 1단계로 사용하고, JEV가 불확실한 문장에서만 LLM을 호출하여 두 점수를 라벨로 "
    "학습한 가중치로 결합하는 캐스케이드를 제안한다."
))

# ---------------- II. 본론 ----------------
heading(doc, "Ⅱ. 본 론")
subheading(doc, "2.1 선행 연구")
body(doc, (
    "NN-PPI[1]는 소형 LLM이 출력한 점수를, 라벨이 있는 보정 세트에서 의미적으로 가까운 이웃 문장들의 잔차로 보정한다. "
    "넘긴 입력에서만 두 모델의 출력을 결합하는 캐스케이드는 일반 분류 과제에서 제안되었다[4]. 반면 본 논문에서는 학습하지 "
    "않은 결정 전용 모델을 1단계로 두고, 그 확률과 LLM 점수의 결합 가중치를 라벨로 학습한다."
))
subheading(doc, "2.2 전체 구조")
body(doc, (
    "그림 1은 제안 방법의 전체 구조를 나타낸다. JEV가 모든 문장에 대해 팩트체크가 필요할 확률 p(x)를 출력하면, 학습 세트에서 "
    "정한 임계값 θ에 가장 가까운 문장, 즉 |logit p(x)−logit θ|가 가장 작은 문장부터 전체의 ρ(호출률)만큼 골라 LLM을 "
    "호출한다. 호출하지 않은 문장은 JEV의 판정(p(x)≥θ)을 그대로 사용하고, 호출한 문장은 2.4절의 결합 판정으로 최종 "
    "판정한다."
))
_dg = pymupdf.open(os.path.join(ROOT, "figures", "jev_method_diagram.pdf"))[0].rect
figure(doc, os.path.join(ROOT, "figures", "jev_method_diagram.png"), _dg.width / 72 * 2.54, "그림 1. JEV 기반 결합형 캐스케이드의 전체 구조")
subheading(doc, "2.3 JEV 판정")
body(doc, (
    "JEV에는 LLM과 같은 NN-PPI의 판정 기준을 구조화된 지시로 전달하고, 문장마다 판정 기준에 따라 팩트체크가 필요한 주장인지를 "
    "묻는 이진 질문으로 확률 p(x)를 얻었다. 질문 문구는 두 후보 중 학습 세트 200문장에서 AUC가 높은 쪽으로 정하였다. JEV는 "
    "추가 학습 없이 사용하며, 라벨은 임계값 θ를 정하는 데만 쓴다."
))
subheading(doc, "2.4 결합 판정")
body(doc, "호출한 문장에 대해 LLM이 출력한 점수 s(x)∈[0,1]와 JEV의 로짓 ℓ(x)=logit p(x)를 다음과 같이 결합한다.")
eq = doc.add_paragraph()
eq.paragraph_format.space_before = Pt(3)
eq.paragraph_format.space_after = Pt(3)
col_w = (21.0 - 2.0 - 1.5 - 0.75) / 2  # one text column (cm)
eq.paragraph_format.tab_stops.add_tab_stop(Cm(col_w / 2), WD_TAB_ALIGNMENT.CENTER)
eq.paragraph_format.tab_stops.add_tab_stop(Cm(col_w), WD_TAB_ALIGNMENT.RIGHT)
eq.add_run("\t")
add_math(eq.add_run(), "eq1")
set_font(eq.add_run("\t(1)"), BODY_FONT, 9)
p = body(doc, "여기서 (a, b, c)는 학습 세트 L에서 학습한 로지스틱 회귀 계수, σ는 시그모이드 함수, ")
p.paragraph_format.widow_control = False
add_math(p.add_run(), "indicator")
set_font(p.add_run(
    "는 조건이 참이면 1인 지시함수이다. 비교 대상인 교체형 캐스케이드(이하 교체형)는 같은 문장에 LLM을 호출하되 ℓ을 "
    "사용하지 않고, s를 L에서 정한 고정 임계값 t와 비교한다."), BODY_FONT, 9)

# ---------------- III. 실험 ----------------
heading(doc, "Ⅲ. 실 험")
subheading(doc, "3.1 실험 환경")
body(doc, (
    f"CLEF 2024 CheckThat! Task 1 영어 데이터[5]는 {R_MAIN[cl]['n_calib_frontier']:,}문장을 학습 세트로, dev·dev-test·공식 "
    f"test를 합친 {R[cl]['n_test']:,}문장을 테스트로 사용하였다. ClaimBuster[6]는 2012년 토론 "
    f"{R_MAIN[cb]['n_calib_frontier']:,}문장과 2016년 토론 {R[cb]['n_test']:,}문장을 각각 학습 세트와 테스트로 사용하였다. "
    "JEV는 TypeSafe API로, LLM은 Claude Sonnet 5로 채점하였으며, LLM은 NN-PPI의 판정 기준 프롬프트로 40문장씩 묶어 "
    "채점하였다. 비교를 위해 JEV 대신 문장 임베딩(all-MiniLM-L6-v2) 기반 RBF-SVM(이하 임베딩 SVM)을 1단계로 쓰는 "
    "결합형과, LLM 전량 호출에 NN-PPI를 적용한 설정도 평가하였다."
))
subheading(doc, "3.2 평가 방법")
body(doc, (
    "평가 지표는 테스트 정확도이며, 학습 세트의 80%를 비복원 추출하여 5회 반복한 평균을 보고한다. 방법 간 차이는 5회 각각의 "
    "예측에 McNemar 검정(α=0.05)을 적용하여 평가하였다. 임계값, 결합 계수, 호출률은 모두 학습 세트 안에서만 정하였다. "
    "호출률은 캐스케이드마다 학습 세트의 교차검증 정확도가 그 이후 0.1%p 이하로만 오르는 가장 작은 값으로 정하였다."
))

subheading(doc, "3.3 실험 결과")


def cell(ds, name):
    """Accuracy in Table 1; † when the proposed row is significantly higher on a majority of the 5 runs."""
    c = f"table:{name}"
    dag = c in R[ds]["tests"] and nsig(ds, c) >= MAJ
    assert not dag or acc(ds, FSEL[ds]) > acc(ds, name)  # † means "lower than the proposed method"
    return f3(acc(ds, name)) + ("†" if dag else "")


def rates(v):
    return f"{pct(rate(cl, v))}/{pct(rate(cb, v))}"


def rule(tc, side, sz):
    """Set one border of a table cell (sz in eighths of a point; 0 removes it)."""
    tcPr = tc._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single" if sz else "nil")
    if sz:
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")
    borders.append(el)


assert all(RSEL[d] == f"jevrep_{rate(d)}" for d in (cl, cb))
rows = [
    ("임베딩 SVM", "0%", cell(cl, "svm"), cell(cb, "svm")),
    ("JEV (임계값 조정)", "0%", cell(cl, "jev_thr"), cell(cb, "jev_thr")),
    ("LLM 전량 호출", "100%", cell(cl, "sonnet_raw"), cell(cb, "sonnet_raw")),
    ("  + 임계값 조정", "100%", cell(cl, "sonnet_thr"), cell(cb, "sonnet_thr")),
    ("  + NN-PPI [1]", "100%", cell(cl, "sonnet_nnppi"), cell(cb, "sonnet_nnppi")),
    ("SVM→LLM 결합형", rates("fuse"), f3(acc(cl, SSEL[cl])) + ("†" if nsig(cl, f"table:{SSEL[cl]}") >= MAJ else ""),
     f3(acc(cb, SSEL[cb])) + ("†" if nsig(cb, f"table:{SSEL[cb]}") >= MAJ else "")),
    ("JEV→LLM 교체형", rates("jevfuse"), f3(acc(cl, RSEL[cl])) + ("†" if nsig(cl, f"table:{RSEL[cl]}") >= MAJ else ""),
     f3(acc(cb, RSEL[cb])) + ("†" if nsig(cb, f"table:{RSEL[cb]}") >= MAJ else "")),
    ("JEV→LLM 결합형(제안)", rates("jevfuse"), f3(acc(cl, FSEL[cl])), f3(acc(cb, FSEL[cb]))),
]
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_before = Pt(4)
cap.paragraph_format.space_after = Pt(2)
set_font(cap.add_run("표 1. 테스트 정확도 (†: 5회 중 3회 이상 제안보다 유의하게 낮음)"), BODY_FONT, 8.5, bold=True)
t = doc.add_table(rows=1 + len(rows), cols=4)
t.style = "Table Grid"  # compact cell paragraphs; its grid lines are switched off below
t.alignment = WD_TABLE_ALIGNMENT.CENTER
_tb = OxmlElement("w:tblBorders")
for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
    _el = OxmlElement(f"w:{side}")
    _el.set(qn("w:val"), "nil")
    _tb.append(_el)
t._tbl.tblPr.append(_tb)
for j, h in enumerate(["방법", "호출률", "CLEF", "ClaimBuster"]):
    c = t.rows[0].cells[j]
    c.text = ""
    set_font(c.paragraphs[0].add_run(h), BODY_FONT, 8.5, bold=True)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for i, (name, rt, a1, a2) in enumerate(rows, start=1):
    bold = "제안" in name
    for j, v in enumerate([name, rt, a1, a2]):
        c = t.rows[i].cells[j]
        c.text = ""
        set_font(c.paragraphs[0].add_run(v), BODY_FONT, 8.5, bold=bold)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
widths = [Cm(3.95), Cm(1.5), Cm(1.15), Cm(1.75)]  # 8.35 cm: one template column
for row in t.rows:
    for c, w in zip(row.cells, widths):
        c.width = w
for c in t.rows[0].cells:
    rule(c, "top", 8)
    rule(c, "bottom", 4)
for c in t.rows[-1].cells:
    rule(c, "top", 4)
    rule(c, "bottom", 8)
cap.paragraph_format.keep_with_next = True
for row in t.rows[:-1]:
    for c in row.cells:
        c.paragraphs[0].paragraph_format.keep_with_next = True

first = body(doc, (
    "표 1은 각 방법의 테스트 정확도를 비교한 결과이다. JEV는 라벨로 학습하지 않았음에도 임베딩 SVM보다 CLEF에서 "
    f"{gain[cl]:.1f}%p, ClaimBuster에서 {gain[cb]:.1f}%p 높았다(5회 모두 유의). JEV 단독의 정확도는 LLM 전량 호출의 더 나은 "
    f"설정과 비교하여 CLEF에서는 5회 중 {nsig(cl, 'jev_thr_vs_best_full')}회 유의하게 낮았으나, ClaimBuster에서는 5회 모두 "
    "유의하게 높았다."
))
first.paragraph_format.space_before = Pt(5)
body(doc, (
    f"제안한 결합형의 호출률은 학습 세트에서 CLEF {pct(rate(cl))}, ClaimBuster {pct(rate(cb))}로 정해졌다. CLEF에서는 "
    f"{pct(rate(cl))} 호출로 {f3(acc(cl, FSEL[cl]))}를 달성하여 LLM 전량 호출({f3(acc(cl, BEST[cl]))})과 5회 모두 유의한 "
    f"차이가 없었고, 임베딩 SVM을 1단계로 쓴 결합형이 {pct(rate(cl, 'fuse'))} 호출로 얻은 정확도({f3(acc(cl, SSEL[cl]))})와도 "
    "차이가 없었다. ClaimBuster에서는 JEV 단독이 이미 LLM 전량 호출보다 높아 학습 세트가 LLM 호출을 선택하지 않았으며, LLM "
    f"호출 없이도 임베딩 SVM 결합형({pct(rate(cb, 'fuse'))} 호출, {f3(acc(cb, SSEL[cb]))})과 차이가 없었다. 다만 테스트에서는 "
    f"{pct(CB_PEAK)} 호출 시 {f3(acc(cb, f'jevfuse_{CB_PEAK}'))}까지 올랐으나, 학습 세트에서는 이득이 나타나지 않아 선택되지 "
    "않았다. 두 데이터셋 모두 NN-PPI를 적용한 전량 호출보다는 5회 모두 유의하게 높았다."
))
body(doc, (
    "그림 2는 호출률에 따른 테스트 정확도를 나타낸다. JEV를 1단계로 쓰면 곡선이 임베딩 SVM보다 높은 곳에서 시작하므로, 같은 "
    "정확도에 훨씬 적은 호출로 도달한다. 결합형은 LLM을 호출하는 모든 호출률에서 교체형보다 평균 정확도가 높았고, 그 차이는 "
    f"CLEF {sig_range(cl)}, ClaimBuster {sig_range(cb)}의 호출률에서 5회 중 3회 이상 유의하였다. 특히 ClaimBuster에서 "
    f"교체형은 호출을 늘릴수록 정확도가 {f3(acc(cb, 'jevrep_0.2'))}에서 {f3(acc(cb, 'jevrep_1.0'))}까지 떨어졌는데, 이는 교체가 "
    "JEV의 옳은 판정을 LLM의 틀린 판정으로 바꾸기 때문으로 해석된다."
))

figure(doc, os.path.join(ROOT, "figures", "jev_budget_full.png"), 8.2,
       "그림 2. 호출률에 따른 테스트 정확도 (5회 평균, 띠는 표준편차)")

# ---------------- IV. 결론 ----------------
heading(doc, "Ⅳ. 결 론")
body(doc, (
    "본 논문에서는 결정 전용 모델 JEV를 1단계로 두고 불확실한 문장에서만 LLM을 호출해 두 점수를 결합하는 캐스케이드를 "
    f"제안하였다. JEV는 라벨로 학습한 분류기보다 정확하여, LLM 호출률을 CLEF에서 {pct(rate(cl))}, ClaimBuster에서 "
    f"{pct(rate(cb))}로 낮추면서도 LLM 전량 호출과 같거나 높은 정확도를 얻었다. 결합은 교체보다 정확도가 높았으나, 두 모델이 "
    "같은 곳에서 틀린다는 보고[3]와 같이 CLEF에서 LLM 전량 호출 대비 정확도 이득은 유의하지 않았다. 다만 JEV는 상용 "
    "서비스이고, 출처가 겹치는[5] 영어 두 데이터셋과 단일 LLM으로 평가하였다. 향후 다국어 데이터와 여러 결정 모델·LLM으로 "
    "평가를 확장할 계획이다."
))

# ACKNOWLEDGMENT omitted (no funding to acknowledge); re-add here if needed

heading(doc, "참 고 문 헌")
refs = [
    "[1] P. Amatya, Venktesh V, and V. Setty, \"Calibrating Small Language Models for Claim Check-Worthiness Detection,\" arXiv:2608.30731, 2026.",
    "[2] Y. Li, Y. Miao, R. Krishnan, and R. Padman, \"JEV-as-a-Judge: Accept When Confident, Escalate When Unsure,\" arXiv:2609.26550, 2026.",
    "[3] D. Rao and C. Callison-Burch, \"JEV vs. LLMs as Rubric Judges: Cheaper, Faster, and Wrong in the Same Places,\" arXiv:2609.29769, 2026.",
    "[4] Y. Zhang et al., \"Calibration-Aware Uncertainty Cascades for Efficient Heterogeneous Model Collaboration,\" arXiv:2609.11446, 2026.",
    "[5] M. Hasanain et al., \"Overview of the CLEF-2024 CheckThat! Lab Task 1 on Check-Worthiness Estimation of Multigenre Content,\" CEUR-WS vol. 3740, pp. 276–286, 2024.",
    "[6] F. Arslan, N. Hassan, C. Li, and M. Tremayne, \"A Benchmark Dataset of Check-Worthy Factual Claims,\" in Proc. ICWSM, vol. 14, pp. 821–829, 2020.",
]
for r in refs:
    ref = body(doc, r, size=9, indent=0, after=0, align=WD_ALIGN_PARAGRAPH.LEFT).paragraph_format
    ref.line_spacing = 1.0
    ref.left_indent, ref.first_line_indent = Pt(12.3), Pt(-12.3)  # hanging indent as in the template

# Word inserts a gap between Hangul and Latin/digits by default ("LLM을" rendered as "LLM 을");
# switch it off everywhere. pPr children must follow the OOXML schema order, so insert before the
# first element that the schema places after autoSpaceDN.
_AFTER_AUTOSPACE = {qn("w:" + n) for n in ("bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
                                            "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
                                            "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")}


def no_autospace(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    anchor = next((c for c in pPr if c.tag in _AFTER_AUTOSPACE), None)
    for tag in ("w:autoSpaceDE", "w:autoSpaceDN"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), "0")
        if anchor is None:
            pPr.append(el)
        else:
            anchor.addprevious(el)


_HANGUL_PAIR = re.compile(r"(?<=[가-힣])(?=[가-힣(])|(?<=[A-Za-z0-9%])(?=[가-힣])")  # also before "(" after Hangul, and between a Latin word or number and its particle (ClaimBuster|에서는), as in HWP


def hangul_breaks(paragraph):
    """Let justified lines break between Hangul syllables, as in HWP, by putting a zero-width space between adjacent
    syllables. Word otherwise breaks Korean only at spaces, so narrow justified columns stretch the gaps between 어절.
    (w:wordWrap=0 would do the same but also splits English words and numbers mid-token.)"""
    ts = [t for t in paragraph._p.iter(qn("w:t")) if t.text]
    for t in ts[:-1]:
        t.text = _HANGUL_PAIR.sub("​", t.text)
    if ts:  # keep the paragraph's last 어절 whole, so no single syllable ("다.") is left alone on the last line
        head, sep, tail = ts[-1].text.rpartition(" ")
        ts[-1].text = _HANGUL_PAIR.sub("​", head) + sep + tail


def no_break_hyphens(paragraph):
    """Replace '-' in text runs by Word's non-breaking hyphen so NN-PPI, RBF-SVM, dev-test never split across lines."""
    for t in list(paragraph._p.iter(qn("w:t"))):
        if "-" not in (t.text or ""):
            continue
        parts, parent, anchor = t.text.split("-"), t.getparent(), t
        t.text = parts[0]
        t.set(qn("xml:space"), "preserve")
        for part in parts[1:]:
            hyph = OxmlElement("w:noBreakHyphen")
            anchor.addnext(hyph)
            nt = OxmlElement("w:t")
            nt.set(qn("xml:space"), "preserve")
            nt.text = part
            hyph.addnext(nt)
            anchor = nt


for para in doc.paragraphs + [q for tbl in doc.tables for row in tbl.rows for cell in row.cells for q in cell.paragraphs]:
    no_autospace(para)
    if para.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
        hangul_breaks(para)
    if not para.text.startswith("["):  # reference entries may break at hyphens
        no_break_hyphens(para)

out = os.path.join(HERE, "jev_kics_draft.docx")
doc.save(out)
print("saved", out)
